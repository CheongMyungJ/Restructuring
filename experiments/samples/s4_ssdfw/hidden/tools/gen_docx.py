#!/usr/bin/env python3
"""generate materials/SSD_FW_아키텍처_설계서_v2.1.docx (Korean, partly outdated on purpose)
usage: gen_docx.py <out.docx>"""
import sys
from docx import Document
from docx.shared import Pt

out = sys.argv[1]
d = Document()
st = d.styles['Normal']
st.font.name = 'Malgun Gothic'
st.font.size = Pt(10)

def h(t, l=1):
    d.add_heading(t, level=l)

def p(t):
    d.add_paragraph(t)

def bl(items):
    for i in items:
        d.add_paragraph(i, style='List Bullet')

def table(rows, header=True):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            t.cell(r, c).text = str(v)
    d.add_paragraph('')

d.add_heading('S4 SSD 컨트롤러 펌웨어 아키텍처 설계서', 0)
p('문서번호: SRK-FW-S4-ARCH   버전: v2.1   작성: FW개발팀')
table([['버전', '일자', '작성자', '내용'],
       ['v1.0', '2016.05', '김지훈', '최초 작성 (hybrid mapping, MLC)'],
       ['v1.5', '2017.03', '김지훈', 'page mapping 전환, journal 도입 (fmt 2.0)'],
       ['v2.0', '2018.02', '최현우', 'write cache, QLC 지원, vendor command'],
       ['v2.1', '2018.09', '박소연', 'SMART 카운터 persist (fmt 3.0), read retry/refresh']])

h('1. 개요')
p('본 문서는 SRK-C400 컨트롤러 기반 S4 NVMe SSD 펌웨어의 구조를 설명한다. 펌웨어는 C99로 작성되며 '
  'ASIC 타겟(ARM Cortex-R5)과 PC 시뮬레이션 타겟(개발/검증용)을 동일 소스로 빌드한다. '
  'HAL(hal/) 아래만 플랫폼별로 다르다.')
bl(['호스트 인터페이스: NVMe 1.3 (namespace 1개, LBA 4KiB)',
    'NAND: TLC(T-gen2) 또는 QLC(Q-gen1), page 16KiB + spare 64B',
    '매핑 단위: 4KiB (page 당 4 slot)',
    'DRAM write cache: 4KiB x 64 entry'])

h('2. 시스템 구성')
table([['구성요소', '내용'],
       ['CPU', 'Cortex-R5 500MHz, TCM 256KB'],
       ['NAND I/F', '4ch x 2way, ONFI 4.0, LDPC ECC (HW)'],
       ['DRAM', 'DDR3 512MB (map table, write cache, PRP bounce)'],
       ['Host', 'PCIe Gen3 x4, NVMe front-end (SQ fetch/CQ post HW)'],
       ['기타', 'TMP-75 온도센서(I2C), PLP 감지(캐패시터 없음)']])
p('펌웨어 계층: host(NVMe 명령) → cache(write cache) → ftl(매핑/GC/WL/BBM/메타) → hal(NAND/timer/temp). '
  'smart, thermal 모듈은 host/ftl 양쪽에서 참조한다.')

h('3. 태스크 모델')
p('S4 펌웨어는 RTOS 없이 단일 main loop 로 동작한다.')
bl(['NVMe FE 인터럽트: SQ entry 를 SRAM ring 에 적재 → main loop 에서 fw_submit() 으로 처리',
    '1ms 타이머 인터럽트: 백그라운드 작업(idle GC, refresh, 온도 샘플링)을 타이머 ISR 에서 직접 수행',
    'PLP 인터럽트: 전원 차단 감지 플래그만 설정 (보조 전원 없음, flush 불가)',
    '명령 처리는 동기식이며 한 번에 하나의 명령만 처리한다.'])
p('※ shutdown: CC.SHN=01b(normal) 시 cache flush + clean checkpoint 기록, 10b(abrupt) 시 '
  'cache flush + journal 기록 (clean mark 없음).')

h('4. FTL 개요')
h('4.1 매핑', 2)
p('LBA(4KiB) → (block, page, slot) page-level 매핑을 DRAM 에 유지한다. 매핑 엔트리는 32bit 이며 '
  'bit31..16 block, bit15..2 page, bit1..0 slot. 0xFFFFFFFF = unmapped, 0xFFFFFFFE = UNC(복구 불가 표시).')
p('사용자 용량은 user pool(system 영역 제외) 전체 slot 수에서 over-provisioning(OP) 비율을 뺀 값을 '
  '8 LBA(32KiB) 단위로 내림하여 결정하며, 포맷 시 superblock 에 기록된다.')
h('4.2 블록 영역', 2)
table([['영역', '블록', '설명'],
       ['Superblock', '0, 1', 'magic/version/geometry/capacity/serial (사본 2개)'],
       ['Meta ring', '2 ~ 9', 'checkpoint + journal, ring 구조'],
       ['User pool', '10 ~ 끝', 'host / GC open block, free block']])
h('4.3 쓰기 경로 / stream', 2)
p('Host stream 과 GC stream 의 2개 open block 을 사용한다. free block 은 할당 시점에 erase 하며 '
  '할당 정책은 erase count 가 가장 낮은 free block 이다. 부분적으로 채워진 page 는 padding 하여 기록한다.')
h('4.4 Garbage Collection', 2)
bl(['Foreground GC: host stream 이 새 block 을 할당할 때 free block 이 4개 미만이면 수행',
    'Victim 선정(GREEDY): valid count 가 가장 작은 used block. 동률이면 가장 오래된 block 을 선택',
    'Idle GC: host idle 2초 이상, free block 8개 미만일 때 tick 마다 1 block',
    'Cost-benefit 정책(CFG_GC_POLICY=COST_BENEFIT)은 실험용으로 코드만 존재'])
h('4.5 Wear Leveling', 2)
p('GC 8회마다 erase count 편차(max-min)를 확인하여 WL_THRESHOLD(20)를 초과하면 erase count 가 '
  '가장 낮은 used block 의 데이터를 이동시켜 해당 block 이 다시 할당되도록 한다 (static WL).')
h('4.6 Bad Block 관리', 2)
bl(['Factory bad: 포맷 시 NAND bad mark 확인, 사용하지 않음',
    'Program fail: 해당 block retire, 실패한 page 는 새 block 에 재기록, 기존 valid data 는 명령 종료 후 이동',
    'Erase fail: retire 후 다른 free block 할당',
    'Available spare 는 남은 spare block 비율로 계산'])
h('4.7 Read Retry / Refresh', 2)
p('ECC fail 시 NAND 타입별 read-retry table 순서대로 재시도한다. RR_REFRESH_LVL 이상의 level 에서 '
  '복구되면 해당 block 을 refresh 대상으로 등록한다. 모든 level 실패 시 Unrecovered Read Error 를 '
  '반환하고 media error 카운터를 증가시킨다. block 별 read count 가 read disturb 한계를 넘으면 '
  'refresh 를 예약하며, refresh 는 fw_tick 에서 1 block 씩 처리한다.')

h('5. 메타데이터 및 전원 차단 복구')
bl(['Checkpoint: part0(header + persist 카운터 + block table) + map part(4096 entry/page)',
    'Journal: map 변경/erase/bad block 기록 + persist 카운터 snapshot. flush/FUA/erase 전에 commit',
    'Checkpoint 주기: host write 4096 LBA 또는 journal 24 page',
    '새 meta block 으로 넘어갈 때 항상 checkpoint 를 먼저 기록',
    'Mount: superblock → meta ring 전체 scan → 최신 완전한 checkpoint → journal replay → 상태 재구성',
    '마지막 page 가 clean checkpoint 가 아니면 unsafe shutdown 으로 카운트'])
p('on-NAND format 은 ftl/meta_format.h 에 정의되어 있으며 fmt 3.x 를 유지해야 한다 (필드 드라이브 호환).')

h('6. Write Cache')
p('4KiB x 64 entry write-back cache. 같은 LBA 재기록 시 entry 를 덮어쓴다. 가득 차면 가장 오래된 '
  '16 entry 를 destage 한다. FLUSH/FUA/shutdown 시 destage 하며, read 는 cache 를 먼저 확인한다. '
  'Volatile Write Cache(feature 0x06) 가 disable 되면 write-through 로 동작한다.')

h('7. SMART / Health')
table([['항목', '계산'],
       ['Composite Temperature', '마지막 샘플 온도(°C) + 273 (Kelvin)'],
       ['Available Spare', '남은 spare block / 초기 spare block x 100'],
       ['Available Spare Threshold', '10%'],
       ['Percentage Used', '평균 erase count x 100 / PE limit, 100 에서 상한'],
       ['Data Units Read/Written', '1000 x 512B 단위 (NVMe spec 에 따름)'],
       ['Power On Hours', 'fw_tick 누적 시간 / 1시간'],
       ['Critical Warning', 'bit0 spare 부족, bit1 온도 임계 초과, bit2 수명 소진']])

h('8. Thermal Throttling')
table([['Level', '진입 온도', '동작'],
       ['0', '-', '정상'],
       ['1', '70°C 이상', 'idle GC 중지'],
       ['2', '80°C 이상', 'write cache 사용량 1/2 제한'],
       ['3', '85°C 이상', 'write cache 미사용 (write-through)']])
p('온도는 1초마다 샘플링하며 내려갈 때는 3°C hysteresis 를 적용한다. Temperature Threshold(feature 0x04) '
  '이상이면 SMART critical warning bit1 을 설정한다.')

h('9. Vendor Command')
p('0xC0 DUMP INTERNAL COUNTERS: 512 byte, u32 필드 (little-endian). GC/WL/bad block/read retry/refresh '
  '카운터, free block 수, erase count 통계, thermal 정보 등. 공장 검사 툴에서 사용한다.')

h('10. 빌드 구성')
table([['옵션', '설명'],
       ['CFG_NAND_TYPE', 'TLC / QLC (geometry, PE limit, read retry table 선택)'],
       ['CFG_OP_PERCENT', 'over-provisioning 비율 (사용자 용량 결정)'],
       ['CFG_GC_POLICY', 'GREEDY / COST_BENEFIT'],
       ['CFG_WL_ENABLE', 'static wear leveling 사용 여부']])
p('기본값은 config/fw_config.h 를 참고한다. 양산 바이너리는 빌드 서버에서 생성한다.')
d.save(out)
print('saved', out)
