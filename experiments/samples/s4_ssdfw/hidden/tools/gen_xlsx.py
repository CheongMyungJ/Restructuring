#!/usr/bin/env python3
"""generate materials/NVMe_명령_지원_명세.xlsx and materials/NAND_프로파일.xlsx
usage: gen_xlsx.py <materials_dir>"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font

outdir = sys.argv[1]


def sheet(wb, title, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in ws.columns:
        w = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(60, max(10, w * 1.3))


wb = Workbook()
sheet(wb, '개요', [
    ['항목', '내용'],
    ['문서', 'S4 NVMe 명령 지원 명세'],
    ['버전', 'v1.4 (2018.11)'],
    ['NVMe 버전', '1.3'],
    ['Namespace', '1개 (NSID=1)'],
    ['LBA Format', '1개 (LBAF0: 4096 byte, metadata 없음)'],
    ['MDTS', '8 (1 MiB)'],
    ['Status 표기', 'SCT/SC (hex)'],
], first=True)
sheet(wb, 'Admin', [
    ['Opcode', '명령', '지원', '비고'],
    ['0x02', 'Get Log Page', 'O', 'LID 0x02(SMART) 만 지원, 그 외 SCT1/SC09 Invalid Log Page'],
    ['0x06', 'Identify', 'O', 'CNS 0(namespace), 1(controller), 2(active NS list). 그 외 Invalid Field'],
    ['0x09', 'Set Features', 'O', 'FID 0x02, 0x04, 0x06. SV=1 이면 SCT1/SC0D'],
    ['0x0A', 'Get Features', 'O', 'SEL 0(current) 1(default) 2(saved=default)'],
    ['0x80', 'Format NVM', 'O', 'LBAF 0 만, SES 0/1 (2=crypto erase 미지원)'],
    ['0x00/0x01/0x04/0x05', 'Delete/Create SQ, CQ', 'X', 'FE HW 가 처리 (FW 미지원 → Invalid Opcode)'],
    ['0x10/0x11', 'FW Commit/Download', 'X', 'factory tool 전용 경로 사용'],
    ['0xC0', 'Vendor: Dump Internal Counters', 'O', 'Vendor 시트 참고'],
])
sheet(wb, 'IO', [
    ['Opcode', '명령', '지원', '비고'],
    ['0x00', 'Flush', 'O', 'NSID 1 또는 0xFFFFFFFF'],
    ['0x01', 'Write', 'O', 'FUA(CDW12 bit30) 지원'],
    ['0x02', 'Read', 'O', ''],
    ['0x08', 'Write Zeroes', 'X', 'Invalid Opcode'],
    ['0x09', 'Dataset Management', 'O', 'AD(deallocate) 만 동작, 나머지 attribute 무시'],
    ['-', '공통 검사 순서', '-', 'NSID → MDTS → LBA range'],
])
sheet(wb, 'Features', [
    ['FID', '이름', '기본값', 'Set 범위 / 비고'],
    ['0x02', 'Power Management', 'PS=0', 'PS 0~2 (NPSS=2), 초과 시 Invalid Field. 저장 안 됨'],
    ['0x04', 'Temperature Threshold', '358 K (85°C)', 'TMPSEL=0, THSEL=0(over) 만 지원'],
    ['0x06', 'Volatile Write Cache', 'Enable(1)', 'disable 시 cache flush 후 write-through'],
    ['-', '공통', '-', '모든 feature 는 전원 인가 시 기본값으로 복귀 (saveable 아님)'],
])
sheet(wb, 'SMART_Log', [
    ['Offset', '필드', '크기', '설명'],
    ['0', 'Critical Warning', '1', 'bit0 spare, bit1 temperature, bit2 reliability'],
    ['1', 'Composite Temperature', '2', 'Kelvin'],
    ['3', 'Available Spare', '1', '%'],
    ['4', 'Available Spare Threshold', '1', '5%'],
    ['5', 'Percentage Used', '1', '%'],
    ['32', 'Data Units Read', '16', '1000 x 512B'],
    ['48', 'Data Units Written', '16', '1000 x 512B'],
    ['64', 'Host Read Commands', '16', ''],
    ['80', 'Host Write Commands', '16', ''],
    ['96', 'Controller Busy Time', '16', '미구현 (0)'],
    ['112', 'Power Cycles', '16', ''],
    ['128', 'Power On Hours', '16', ''],
    ['144', 'Unsafe Shutdowns', '16', ''],
    ['160', 'Media and Data Integrity Errors', '16', ''],
    ['176', 'Number of Error Information Log Entries', '16', ''],
    ['200', 'Temperature Sensor 1', '2', 'Kelvin'],
])
sheet(wb, 'Vendor', [
    ['Opcode', '명령', '데이터', '설명'],
    ['0xC0', 'Dump Internal Counters', '512 byte', 'u32 필드 23개, CDW0 = 필드 수'],
    ['', 'offset 0x00', 'signature', "'S4DC'"],
    ['', 'offset 0x04', 'layout version', '3'],
    ['', 'offset 0x08', 'GC count', ''],
    ['', 'offset 0x0C', 'WL count', ''],
    ['', 'offset 0x10', 'bad block total', 'factory + grown'],
    ['', 'offset 0x14', 'grown bad blocks', ''],
    ['', 'offset 0x18', 'read retry count', 'retry 시도 횟수'],
    ['', 'offset 0x1C', 'refresh count', ''],
    ['', 'offset 0x20', 'refresh (read retry)', ''],
    ['', 'offset 0x24', 'refresh (read disturb)', ''],
    ['', 'offset 0x28', 'free blocks', ''],
    ['', 'offset 0x2C/0x30/0x34', 'min/max/avg erase count', ''],
    ['', 'offset 0x38/0x3C', 'program/erase fail', ''],
    ['', 'offset 0x40/0x44', 'thermal level / throttle events', ''],
    ['', 'offset 0x48', 'cache dirty entries', ''],
    ['', 'offset 0x4C', 'checkpoint count', ''],
    ['', 'offset 0x50', 'NAND data page programs', ''],
    ['', 'offset 0x54', 'format count', ''],
    ['', 'offset 0x58', 'debug level', ''],
])
wb.save(outdir + '/NVMe_명령_지원_명세.xlsx')

wb = Workbook()
sheet(wb, 'Geometry', [
    ['항목', 'TLC (T-gen2)', 'QLC (Q-gen1)', '비고'],
    ['Page size', 16384, 16384, 'byte'],
    ['Spare (FW 사용)', 64, 64, 'byte'],
    ['Pages / block', 48, 64, 'TLC: 16 WL x 3, QLC: 16 WL x 4'],
    ['Blocks', 96, 64, 'sim/EVB 구성 기준'],
    ['PE limit', 3000, 1000, 'SMART percentage used 기준'],
    ['Read disturb limit', 20000, 8000, 'block 당 read 횟수, 초과 시 refresh'],
    ['RR refresh level', 4, 6, '이 level 이상에서 복구되면 refresh'],
    ['Factory bad 표시', 'page0 spare[0] != 0xFF', 'page0 spare[0] != 0xFF', ''],
], first=True)
sheet(wb, 'ReadRetry_TLC', [
    ['순서', 'Retry option (feature 0x89)', '비고'],
    *[[i + 1, lv, ''] for i, lv in enumerate([1, 2, 3, 4, 5, 6, 7])],
])
sheet(wb, 'ReadRetry_QLC', [
    ['순서', 'Retry option (feature 0x89)', '비고'],
    *[[i + 1, lv, 'vendor 권고 순서 (2018.06)' if i == 0 else ''] for i, lv in
      enumerate([2, 4, 6, 8, 10, 12, 14, 1, 3, 5])],
])
sheet(wb, '이력', [
    ['일자', '내용'],
    ['2016.03', 'TLC T-gen2 profile'],
    ['2017.08', 'QLC Q-gen1 샘플 profile 추가 (EVB)'],
    ['2018.06', 'QLC read retry 순서 변경'],
])
wb.save(outdir + '/NAND_프로파일.xlsx')
print('saved')
