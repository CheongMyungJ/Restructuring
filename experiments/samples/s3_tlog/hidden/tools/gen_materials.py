#!/usr/bin/env python3
"""
Generate the hand-over materials (Korean) into ../../materials:
  TLOG_배치_설계서_v1.3.docx   (python-docx)  - partially OUTDATED on purpose
  인터페이스_정의서.xlsx        (openpyxl)     - one error code (E08) left undocumented
  운영_가이드.txt
Nothing here may reveal production JAVA_OPTS / time zone / duplicate policy / truncation intent /
PURGE being unused.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).resolve().parents[2] / "materials"
FONT = "맑은 고딕"


# =============================================================================================
# 설계서 (docx)
# =============================================================================================
def set_font(run, size=None, bold=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def para(doc, text, size=10, bold=False, align=None, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size, bold)
    if align:
        p.alignment = align
    return p


def heading(doc, text, level):
    h = doc.add_heading("", level=level)
    r = h.add_run(text)
    set_font(r, {1: 15, 2: 12, 3: 11}[level], True)
    return h


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        set_font(c.paragraphs[0].add_run(h), 9, True)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            set_font(cells[i].paragraphs[0].add_run(str(v)), 9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Pt(w)
    doc.add_paragraph()
    return t


def bullets(doc, items, size=10):
    for it in items:
        para(doc, it, size, style="List Bullet")


def design_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(10)

    for _ in range(4):
        doc.add_paragraph()
    para(doc, "TLOG 배치 설계서", 24, True, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "SSD 양산 테스트 로그 수집·판정·집계 배치", 14, False, WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    para(doc, "문서번호 : MIT-TLOG-DS-001", 11, False, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "버전 : v1.3", 11, False, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "작성일 : 2016.12.20", 11, False, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "작성 : 제조IT팀 박선영 / 검토 : 김정훈 (PL)", 11, False, WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(6):
        doc.add_paragraph()
    para(doc, "DH반도체 제조IT팀", 12, True, WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    heading(doc, "개정 이력", 1)
    table(doc, ["버전", "일자", "변경 내용", "작성자"], [
        ["v1.0", "2015.09.21", "최초 작성 (ADV 테스터 로그 적재/판정/수율)", "김정훈"],
        ["v1.1", "2016.03.30", "BIN 룰 우선순위 반영, 재테스트 인정 기준 추가", "박선영"],
        ["v1.2", "2016.11.25", "TER 테스터 추가(CSV, UTC), ES 샘플 수율 제외", "이동욱"],
        ["v1.3", "2016.12.20", "수율 경보(ALERT) 배치 추가, 경보 이력 테이블 추가", "박선영"],
    ], [40, 60, 260, 50])

    heading(doc, "1. 개요", 1)
    heading(doc, "1.1 목적", 2)
    para(doc, "본 문서는 SSD 양산 테스트 공정(FT1 / BI / FT2)에서 테스터 장비가 생성하는 로그 파일을 수집하여 "
              "디바이스별 최종 BIN 을 판정하고, LOT 별 수율을 집계하여 MES 로 전달하는 TLOG 배치의 설계 내용을 기술한다.")
    heading(doc, "1.2 범위", 2)
    bullets(doc, ["배치 : LOAD, JUDGE, YIELD, ALERT, PURGE",
                  "연계 파일 : 장비 로그(ADV/TER), MES 수율 리포트, 수율 경보 파일, reject 파일",
                  "상세 파일 레이아웃 및 에러코드는 별도 문서 '인터페이스 정의서' 참조"])
    heading(doc, "1.3 용어", 2)
    table(doc, ["용어", "설명"], [
        ["시리얼(SERIAL)", "디바이스 고유 번호"],
        ["LOT", "생산 투입 단위 묶음"],
        ["STEP", "테스트 공정 구분. FT1(1차 기능), BI(Burn-In), FT2(2차 기능)"],
        ["BIN", "디바이스 최종 판정 등급 (1~9). 양품 여부는 BIN 코드의 PASS_YN 으로 관리"],
        ["근무일자(WORK_DT)", "06시 ~ 익일 06시를 하나의 생산일로 보는 일자"],
        ["ES 샘플", "엔지니어링 샘플. 시리얼이 'ES' 로 시작"],
        ["기준일자(BIZ_DATE)", "배치 실행 대상 일자 (통상 전일)"],
        ["MES", "제조실행시스템. TLOG 수율 리포트를 수거하여 LOT 이력에 반영"],
    ], [110, 300])

    heading(doc, "2. 시스템 구성", 1)
    heading(doc, "2.1 구성도", 2)
    para(doc, "[ADV 테스터] --(FTP, *.dat)-->  /data/tlog/inbox\n"
              "[TER 테스터] --(SFTP, *.csv)--> /data/tlog/inbox\n"
              "배치서버(tlogbat01) : TLOG 배치 (cron 실행)  <-->  TLOGDB (Oracle 11g)\n"
              "/data/tlog/outbox  --> MES (07:30 수거), 품질팀 경보 메일\n"
              "/data/tlog/reject  --> 장비 업체 회신\n"
              "/data/tlog/archive --> 처리 완료 파일 보관", 9)
    heading(doc, "2.2 S/W 구성", 2)
    table(doc, ["구분", "내용"], [
        ["OS", "CentOS 6.7"],
        ["Java", "JDK 1.7"],
        ["Framework", "Spring Framework 4.1, Spring Batch 3.0"],
        ["Persistence", "MyBatis 3.2 (일부 JdbcTemplate)"],
        ["DB", "Oracle 11g (TLOGDB)"],
        ["실행", "CommandLineJobRunner + Shell (cron)"],
    ], [110, 300])
    heading(doc, "2.3 디렉토리", 2)
    table(doc, ["경로", "용도"], [
        ["/app/tlog/batch", "배치 프로그램 (bin, lib)"],
        ["/app/tlog/logs", "실행 로그"],
        ["/data/tlog/inbox", "장비 로그 수신"],
        ["/data/tlog/outbox", "MES 리포트, 수율 경보 파일"],
        ["/data/tlog/reject", "오류 레코드 파일 (*.rej)"],
        ["/data/tlog/archive/<기준일자>", "처리 완료 파일"],
    ], [160, 250])

    heading(doc, "3. 배치 목록", 1)
    table(doc, ["Job ID", "Job 명", "실행 주기", "선행", "설명"], [
        ["LOAD", "tlogLoadJob", "매일 06:30", "-", "장비 로그 파일 적재"],
        ["JUDGE", "binJudgeJob", "매일 06:40", "LOAD", "디바이스별 최종 BIN 판정"],
        ["YIELD", "yieldSummaryJob", "매일 06:50", "JUDGE", "LOT 수율 집계, MES 리포트 생성"],
        ["ALERT", "yieldAlertJob", "매일 07:00", "YIELD", "수율 경보 파일 생성, 경보 이력 저장"],
        ["PURGE", "purgeJob", "매주 일요일 03:00", "-", "보관기간 경과 테스트 결과 삭제"],
    ], [45, 85, 75, 45, 160])
    para(doc, "※ 기준일자(BIZ_DATE)는 전일(D-1) 이며 실행 스크립트 인자로 전달한다.", 9)

    heading(doc, "4. 배치 상세", 1)
    heading(doc, "4.1 LOAD (tlogLoadJob)", 2)
    para(doc, "(1) 대상 파일", 10, True)
    bullets(doc, ["inbox 디렉토리에서 파일명의 일자가 기준일자와 같은 파일",
                  "ADV : ADV_yyyyMMdd_NN.dat (고정길이, MS949)",
                  "TER : TER_yyyyMMdd_NN.csv (CSV, UTF-8)",
                  "파일명 규칙에 맞지 않는 파일은 처리하지 않고 inbox 에 그대로 둔다."])
    para(doc, "(2) 처리 흐름", 10, True)
    bullets(doc, ["① 대상 파일 목록 조회",
                  "② 트레일러 건수 검증 - 불일치 시 해당 파일 전체를 reject (적재하지 않음)",
                  "③ 레코드 파싱 및 검증 - 오류 레코드는 reject 파일에 기록하고 계속 진행",
                  "④ 중복 체크",
                  "⑤ TB_TEST_RESULT 저장",
                  "⑥ 파일별 처리 결과를 TB_FILE_LOG 에 기록 (전체/적재/오류/중복 건수)",
                  "⑦ 처리 완료 파일을 archive/<기준일자> 로 이동"])
    para(doc, "(3) 시각 및 근무일자", 10, True)
    bullets(doc, ["ADV 테스터는 공장 현지 시각으로 기록한다.",
                  "TER 테스터는 UTC 로 기록하므로 KST(UTC+9) 로 변환하여 저장한다.",
                  "근무일자(WORK_DT)는 시작일시 기준으로 산출하며, 06시 이전에 시작한 테스트는 전일 근무로 본다."])
    para(doc, "(4) 데이터 정규화", 10, True)
    bullets(doc, ["시리얼 : 공백 제거, 대문자 변환, 하이픈(-) 제거. TER 는 'SN' 접두어가 붙어 오는 경우가 있어 제거한다.",
                  "결과 : TER 의 PASS/FAIL/ABORT 를 P/F/A 로 변환",
                  "온도 : TER 는 화씨(℉)로 기록되므로 섭씨로 변환 후 반올림하여 정수로 저장"])
    para(doc, "(5) 중복 처리", 10, True)
    para(doc, "동일 시리얼 + STEP + 시작일시 데이터가 이미 존재하면 기존 데이터를 유지하고, 신규 건은 적재하지 않고 "
              "중복 건수(DUP_CNT)로 집계한다.")
    para(doc, "(6) 오류 처리", 10, True)
    para(doc, "검증 오류 코드 및 reject 파일 형식은 인터페이스 정의서를 따른다. reject 파일은 장비 업체 회신용으로 사용한다.")

    heading(doc, "4.2 JUDGE (binJudgeJob)", 2)
    para(doc, "기준일자와 근무일자(WORK_DT)가 같은 테스트 결과를 대상으로 시리얼 단위로 최종 BIN 을 판정한다.")
    para(doc, "(1) 유효 시도", 10, True)
    para(doc, "STEP 별 재테스트는 시작일시 순으로 최대 5회까지 인정하며, 인정된 시도 중 마지막 시도를 해당 STEP 의 "
              "유효 결과로 한다. (2016.01 공정기술팀 협의)")
    para(doc, "(2) 판정 순서", 10, True)
    bullets(doc, ["① 이전 근무일에 폐기(BIN 9) 판정 이력이 있는 디바이스는 BIN 9 유지",
                  "② 유효 시도 중 ABORT 가 있으면 BIN 8",
                  "③ BIN 룰(TB_BIN_RULE, USE_YN='Y') 적용 - 첫 번째로 조건을 만족하는 룰의 BIN",
                  "④ 해당 룰이 없으면 모든 유효 시도가 PASS 인 경우 BIN 1, 그 외 BIN 7"])
    para(doc, "(3) BIN 룰", 10, True)
    bullets(doc, ["적용 대상 : 제품코드가 일치하거나 '*' (전 제품)",
                  "적용 순서 : PRIORITY 오름차순 → 동일 우선순위는 제품 지정 룰 우선 → BIN 번호가 큰 룰 우선",
                  "판정 항목 : READ_MBPS, WRITE_MBPS (유효 시도 중 최소값) / LAT_US, BAD_BLK, ECC_CNT, TEMP_MAX (최대값)",
                  "비교 연산자 : <, <=, >, >=",
                  "제품별 예외 기준은 반드시 BIN 룰 테이블에 등록하여 관리한다. (프로그램 하드코딩 금지)"])
    para(doc, "(4) ES 샘플", 10, True)
    para(doc, "엔지니어링 샘플(시리얼 'ES' 시작)은 판정은 수행하되 수율 집계 제외 대상으로 표시(EXCL_YN='Y')한다.")
    para(doc, "(5) 저장", 10, True)
    para(doc, "TB_DEVICE_BIN 에 (시리얼, 근무일자) 기준으로 MERGE 한다. 재실행 시 결과를 덮어쓴다.")

    heading(doc, "4.3 YIELD (yieldSummaryJob)", 2)
    bullets(doc, ["TB_DEVICE_BIN (EXCL_YN='N') 을 제품/LOT 별로 집계하여 TB_LOT_YIELD 에 저장",
                  "투입수량 = 디바이스 수, 양품수량 = PASS_YN='Y' 인 BIN 의 디바이스 수, BIN1~BIN9 별 수량",
                  "수율(%) = 양품수량 / 투입수량 × 100 (소수점 셋째 자리에서 반올림하여 둘째 자리까지)",
                  "해당 근무일 데이터를 삭제 후 재생성 (재실행 가능)",
                  "MES 리포트 파일 생성 : outbox/YIELD_yyyyMMdd.txt (레이아웃은 인터페이스 정의서 참조)"])

    heading(doc, "4.4 ALERT (yieldAlertJob)", 2)
    bullets(doc, ["TB_LOT_YIELD 중 수율이 제품별 경보 기준(TB_PRODUCT.ALERT_YIELD) 미만인 LOT 추출",
                  "경보 이력을 TB_YIELD_ALERT 에 저장",
                  "outbox/ALERT_yyyyMMdd.csv 생성 → 품질팀 경보 메일 스크립트에서 사용",
                  "경보 대상이 없으면 파일을 생성하지 않는다."])

    heading(doc, "4.5 PURGE (purgeJob)", 2)
    para(doc, "TB_SYS_CONFIG 의 RETENTION_DAYS (기본 90일) 이전 근무일자의 TB_TEST_RESULT 데이터를 삭제한다. "
              "매주 일요일 03:00 실행.")

    heading(doc, "5. 데이터 흐름", 1)
    table(doc, ["단계", "입력", "처리", "출력"], [
        ["1", "inbox/*.dat, *.csv", "LOAD", "TB_TEST_RESULT, TB_FILE_LOG, reject/*.rej, archive/"],
        ["2", "TB_TEST_RESULT", "JUDGE", "TB_DEVICE_BIN"],
        ["3", "TB_DEVICE_BIN", "YIELD", "TB_LOT_YIELD, outbox/YIELD_*.txt"],
        ["4", "TB_LOT_YIELD", "ALERT", "TB_YIELD_ALERT, outbox/ALERT_*.csv"],
    ], [35, 120, 55, 200])

    heading(doc, "6. 테이블 목록", 1)
    table(doc, ["테이블", "설명", "관리"], [
        ["TB_PRODUCT", "제품 기준정보 (제품명, 경보 기준 수율, 사용여부)", "기준정보팀"],
        ["TB_BIN_CODE", "BIN 코드 (양품 여부)", "기준정보팀"],
        ["TB_BIN_RULE", "BIN 판정 룰", "기준정보팀"],
        ["TB_SYS_CONFIG", "시스템 설정", "제조IT팀"],
        ["TB_TEST_RESULT", "테스트 결과", "TLOG"],
        ["TB_FILE_LOG", "파일 적재 이력", "TLOG"],
        ["TB_DEVICE_BIN", "디바이스 최종 BIN", "TLOG"],
        ["TB_LOT_YIELD", "LOT 수율", "TLOG"],
        ["TB_YIELD_ALERT", "수율 경보 이력", "TLOG"],
    ], [100, 240, 70])

    heading(doc, "7. 운영", 1)
    bullets(doc, ["실행 : /app/tlog/batch/bin/run_job.sh <JOB_ID> <기준일자 yyyyMMdd>",
                  "종료코드 0 = 정상, 그 외 = 오류",
                  "동일 기준일자 재실행 시 Job 인스턴스 중복 오류가 발생하므로 BATCH_JOB_* 메타 테이블의 해당 실행 이력을 정리한 후 재실행",
                  "LOAD 재처리 : archive 의 파일을 inbox 로 복사한 후 LOAD 재실행",
                  "기준정보 변경은 기준정보팀에 '기준정보 변경요청서'로 요청"])

    heading(doc, "8. BIN 코드 (참고)", 1)
    para(doc, "BIN 코드는 기준정보팀에서 관리하며 아래는 작성 시점 기준이다.", 9)
    table(doc, ["BIN", "의미", "양품"], [
        ["1", "양품", "Y"], ["2", "양품 (성능 하위 등급)", "Y"], ["3", "READ 성능 미달", "N"],
        ["4", "WRITE 성능 미달", "N"], ["5", "불량 블록 과다", "N"], ["6", "온도 이상", "N"],
        ["7", "기능 불량", "N"], ["8", "테스트 중단", "N"], ["9", "폐기", "N"],
    ], [40, 250, 50])
    doc.save(str(OUT / "TLOG_배치_설계서_v1.3.docx"))


# =============================================================================================
# 인터페이스 정의서 (xlsx)
# =============================================================================================
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFILL = PatternFill("solid", fgColor="D9E1F2")
TFONT = Font(name=FONT, size=12, bold=True)
NFONT = Font(name=FONT, size=10)
BFONT = Font(name=FONT, size=10, bold=True)


def sheet(wb, title, heading_lines, header, rows, widths):
    ws = wb.create_sheet(title)
    r = 1
    ws.cell(r, 1, heading_lines[0]).font = TFONT
    r += 1
    for line in heading_lines[1:]:
        ws.cell(r, 1, line).font = NFONT
        r += 1
    r += 1
    for i, h in enumerate(header, 1):
        c = ws.cell(r, i, h)
        c.font, c.fill, c.border = BFONT, HFILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        r += 1
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v)
            c.font, c.border = NFONT, BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    return ws


def interface_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "표지"
    ws["A1"] = "TLOG 인터페이스 정의서"
    ws["A1"].font = Font(name=FONT, size=16, bold=True)
    rows = [("문서번호", "MIT-TLOG-IF-002"), ("시스템", "TLOG 배치 (SSD 테스트 로그)"), ("작성", "제조IT팀"),
            ("최종수정", "2017.04.18 (YIELD 리포트 경보여부 추가)")]
    for i, (k, v) in enumerate(rows, 3):
        ws.cell(i, 1, k).font = BFONT
        ws.cell(i, 2, v).font = NFONT
    ws.cell(8, 1, "개정이력").font = BFONT
    hist = [("v1.0", "2015.09.15", "ADV 로그, MES 리포트 최초 정의"),
            ("v1.1", "2016.11.20", "TER 로그(CSV) 추가"),
            ("v1.2", "2016.12.15", "ALERT 파일 추가"),
            ("v1.3", "2017.04.18", "YIELD 리포트 D 레코드에 경보여부(1) 추가, 제품명 22->20 byte")]
    for i, h in enumerate(("버전", "일자", "내용"), 1):
        c = ws.cell(9, i, h)
        c.font, c.fill, c.border = BFONT, HFILL, BORDER
    for r, row in enumerate(hist, 10):
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v)
            c.font, c.border = NFONT, BORDER
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60

    lay = ["레코드", "No", "항목명", "항목ID", "시작위치", "길이(byte)", "타입", "설명"]
    adv = [
        ["H", 1, "레코드구분", "REC_TYPE", 1, 1, "X", "'H'"],
        ["H", 2, "파일생성일자", "FILE_DT", 2, 8, "9", "yyyyMMdd"],
        ["H", 3, "파일생성시각", "FILE_TM", 10, 6, "9", "HHmmss"],
        ["H", 4, "장비업체", "VENDOR", 16, 10, "X", "'ADV'"],
        ["H", 5, "사이트", "SITE", 26, 4, "X", ""],
        ["H", 6, "레이아웃버전", "LAYOUT_VER", 30, 10, "X", "'LAYOUT-V2'"],
        ["H", 7, "FILLER", "", 40, 101, "X", "공백"],
        ["D", 1, "레코드구분", "REC_TYPE", 1, 1, "X", "'D'"],
        ["D", 2, "시리얼", "SERIAL", 2, 16, "X", "좌측정렬, 우측 공백"],
        ["D", 3, "LOT ID", "LOT_ID", 18, 10, "X", ""],
        ["D", 4, "제품코드", "PROD_CD", 28, 8, "X", ""],
        ["D", 5, "STEP", "STEP", 36, 4, "X", "'FT1 ', 'BI  ', 'FT2 '"],
        ["D", 6, "시작일시", "START_DTM", 40, 14, "9", "yyyyMMddHHmmss (공장 현지시각)"],
        ["D", 7, "종료일시", "END_DTM", 54, 14, "9", "yyyyMMddHHmmss (공장 현지시각)"],
        ["D", 8, "테스터ID", "TESTER_ID", 68, 8, "X", ""],
        ["D", 9, "작업자명", "OPER_NM", 76, 10, "X", "한글 1자 = 2byte"],
        ["D", 10, "결과", "RESULT", 86, 2, "X", "'P ', 'F ', 'A '"],
        ["D", 11, "READ 속도", "READ_MBPS", 88, 6, "9(5)V9", "MB/s, 소수 1자리 내포 (005123 = 512.3)"],
        ["D", 12, "WRITE 속도", "WRITE_MBPS", 94, 6, "9(5)V9", "MB/s, 소수 1자리 내포"],
        ["D", 13, "지연시간", "LAT_US", 100, 6, "9", "us"],
        ["D", 14, "불량블록수", "BAD_BLK", 106, 5, "9", ""],
        ["D", 15, "ECC 정정수", "ECC_CNT", 111, 8, "9", ""],
        ["D", 16, "최고온도", "MAX_TEMP", 119, 4, "S9(3)", "섭씨, 부호 포함 ('+072', '-005')"],
        ["D", 17, "FW 버전", "FW_VER", 123, 8, "X", ""],
        ["D", 18, "FILLER", "", 131, 10, "X", "공백"],
        ["T", 1, "레코드구분", "REC_TYPE", 1, 1, "X", "'T'"],
        ["T", 2, "데이터건수", "DETAIL_CNT", 2, 7, "9", "D 레코드 건수"],
        ["T", 3, "FILLER", "", 9, 132, "X", "공백"],
    ]
    sheet(wb, "ADV_레이아웃", ["ADV 테스터 로그 레이아웃",
                              "파일명 : ADV_yyyyMMdd_NN.dat  (yyyyMMdd = 파일 생성일, NN = 일련번호 2자리)",
                              "인코딩 : MS949 / 레코드 길이 : 140 byte 고정 (줄바꿈 CRLF 별도)",
                              "구성 : H 레코드 1건 + D 레코드 n건 + T 레코드 1건",
                              "시작위치는 1부터 시작하는 byte 위치"], lay, adv, [8, 5, 14, 12, 9, 10, 9, 42])

    ter_h = ["No", "컬럼명", "타입", "필수", "설명", "예시"]
    ter = [
        [1, "SERIAL", "문자", "Y", "디바이스 시리얼 ('SN' 접두어, 하이픈 포함되어 오는 경우 있음)", "SN-K1611250001"],
        [2, "LOT", "문자", "Y", "LOT ID", "L16112501"],
        [3, "PRODUCT", "문자", "Y", "제품코드", "C256A110"],
        [4, "STEP", "문자", "Y", "FT1 / BI / FT2", "FT1"],
        [5, "START_UTC", "일시", "Y", "시작일시 ISO-8601 UTC", "2016-11-24T23:10:00Z"],
        [6, "END_UTC", "일시", "Y", "종료일시 ISO-8601 UTC", "2016-11-24T23:25:00Z"],
        [7, "TESTER", "문자", "Y", "테스터 ID", "TER-07"],
        [8, "OPERATOR", "문자", "N", "작업자명 (콤마 포함 시 쌍따옴표)", "\"Kim, Minsu\""],
        [9, "RESULT", "문자", "Y", "PASS / FAIL / ABORT", "PASS"],
        [10, "READ_MBPS", "숫자", "Y", "MB/s", "512.3"],
        [11, "WRITE_MBPS", "숫자", "Y", "MB/s", "498.0"],
        [12, "LAT_US", "정수", "Y", "us", "91"],
        [13, "BAD_BLK", "정수", "Y", "불량 블록 수", "3"],
        [14, "ECC_CNT", "정수", "Y", "ECC 정정 수", "520"],
        [15, "TEMP_F", "숫자", "Y", "최고온도 (화씨)", "150.8"],
        [16, "FW", "문자", "Y", "FW 버전", "SA16A02"],
    ]
    sheet(wb, "TER_컬럼", ["TER 테스터 로그 (CSV)",
                          "파일명 : TER_yyyyMMdd_NN.csv",
                          "인코딩 : UTF-8 (BOM 이 포함될 수 있음) / 구분자 : 콤마(,) / 첫 행은 헤더",
                          "값에 콤마가 포함되면 쌍따옴표로 감싼다",
                          "마지막 행에 '#EOF,<데이터건수>' 트레일러가 올 수 있음 (있는 경우 건수 검증)"],
          ter_h, ter, [5, 14, 8, 6, 52, 26])

    rpt = [
        ["H", 1, "레코드구분", "", 1, 1, "X", "'H'"],
        ["H", 2, "기준일자", "", 2, 8, "9", "yyyyMMdd"],
        ["H", 3, "시스템명", "", 10, 10, "X", "'TLOGBATCH' + 공백"],
        ["H", 4, "데이터건수", "", 20, 7, "9", "D 레코드 건수, 앞자리 0"],
        ["H", 5, "FILLER", "", 27, 34, "X", "공백"],
        ["D", 1, "레코드구분", "", 1, 1, "X", "'D'"],
        ["D", 2, "제품코드", "PROD_CD", 2, 8, "X", "좌측정렬, 공백채움"],
        ["D", 3, "LOT ID", "LOT_ID", 10, 10, "X", "좌측정렬, 공백채움"],
        ["D", 4, "제품명", "PROD_NM", 20, 20, "X", "한글 2byte. 초과분 절삭, 부족분 공백채움"],
        ["D", 5, "투입수량", "IN_QTY", 40, 7, "9", "앞자리 0"],
        ["D", 6, "양품수량", "PASS_QTY", 47, 7, "9", "앞자리 0"],
        ["D", 7, "수율", "YIELD_PCT", 54, 6, "999.99", "소수점 2자리 (예: 099.50)"],
        ["D", 8, "경보여부", "ALERT_YN", 60, 1, "X", "Y / N"],
        ["T", 1, "레코드구분", "", 1, 1, "X", "'T'"],
        ["T", 2, "데이터건수", "", 2, 7, "9", "D 레코드 건수"],
        ["T", 3, "양품수량합계", "", 9, 7, "9", "D 레코드 양품수량 합 (7자리 초과 시 하위 7자리)"],
        ["T", 4, "FILLER", "", 16, 45, "X", "공백"],
    ]
    sheet(wb, "YIELD_리포트", ["MES 수율 리포트",
                              "파일명 : outbox/YIELD_yyyyMMdd.txt (yyyyMMdd = 기준일자)",
                              "인코딩 : MS949 / 줄바꿈 : CRLF / 레코드 길이 : 60 byte 고정",
                              "D 레코드 정렬 : 제품코드, LOT ID 오름차순",
                              "MES 가 매일 07:30 수거"], lay, rpt, [8, 5, 14, 12, 9, 10, 9, 42])

    alert = [
        [1, "BIZ_DATE", "기준일자 yyyyMMdd"],
        [2, "PROD_CD", "제품코드"],
        [3, "PROD_NM", "제품명 (쌍따옴표로 감쌈)"],
        [4, "LOT_ID", "LOT ID"],
        [5, "IN_QTY", "투입수량"],
        [6, "PASS_QTY", "양품수량"],
        [7, "YIELD_PCT", "수율 (소수점 2자리)"],
        [8, "ALERT_YIELD", "제품 경보 기준 수율 (소수점 2자리)"],
    ]
    sheet(wb, "ALERT_CSV", ["수율 경보 파일",
                           "파일명 : outbox/ALERT_yyyyMMdd.csv",
                           "인코딩 : UTF-8 / 줄바꿈 : LF / 첫 행 헤더",
                           "대상 : 수율이 제품 경보 기준 미만인 LOT (정렬: 제품코드, LOT ID)",
                           "경보 대상이 없으면 파일을 생성하지 않음"],
          ["No", "컬럼명", "설명"], alert, [5, 16, 50])

    rej = [
        [1, "라인번호", "원본 파일의 라인 번호 (헤더 포함, 1부터). 파일 단위 오류는 0"],
        [2, "에러코드", "에러코드 시트 참조"],
        [3, "원본레코드", "원본 라인 그대로 (파일 단위 오류는 파일명)"],
    ]
    sheet(wb, "REJECT_파일", ["오류(reject) 파일",
                             "파일명 : reject/<원본파일명>.rej  (예: ADV_20161125_01.dat.rej)",
                             "형식 : 라인번호|에러코드|원본레코드  (구분자 '|')",
                             "인코딩 : 원본 파일과 동일 (ADV = MS949, TER = UTF-8) / 줄바꿈 : LF",
                             "오류 레코드가 있을 때만 생성"],
          ["No", "항목", "설명"], rej, [5, 14, 60])

    err = [
        ["E01", "레코드 길이 오류(ADV) / 컬럼 수 오류(TER)", "레코드"],
        ["E02", "시리얼 누락", "레코드"],
        ["E03", "일시 형식 오류", "레코드"],
        ["E04", "종료일시가 시작일시보다 이전", "레코드"],
        ["E05", "결과값 오류", "레코드"],
        ["E06", "숫자 형식 오류", "레코드"],
        ["E07", "미등록 제품 또는 미사용 제품", "레코드"],
        ["E09", "트레일러 건수 불일치 (파일 전체 미적재)", "파일"],
    ]
    sheet(wb, "에러코드", ["검증 에러코드"], ["코드", "내용", "단위"], err, [8, 50, 10])
    wb.save(str(OUT / "인터페이스_정의서.xlsx"))


# =============================================================================================
# 운영 가이드 (txt)
# =============================================================================================
OPS = r"""==============================================================================
 TLOG 배치 운영 가이드
 (SSD 양산 테스트 로그 수집 / BIN 판정 / LOT 수율 집계)

 최종수정 : 2019.11.28  제조IT팀 이동욱
==============================================================================

1. 개요
 - 테스터(ADV, TER) 로그 파일을 수집하여 디바이스 BIN 을 판정하고
   LOT 수율 리포트를 MES 로 전달하는 배치
 - 배치서버 : tlogbat01 (운영) / tlogbat-dev01 (개발)
 - 실행계정 : tlogadm

2. 디렉토리
 /app/tlog/batch          배치 프로그램 (빌드한 target 디렉토리를 통째로 배포)
 /app/tlog/batch/bin      실행 스크립트 run_job.sh
 /app/tlog/logs           cron 실행 로그
 $TLOG_HOME (/data/tlog)
   inbox/                 장비 로그 수신 (ADV: FTP, TER: SFTP)
   outbox/                MES 리포트(YIELD_*.txt), 수율 경보(ALERT_*.csv)
   reject/                오류 레코드 (*.rej) - 장비 업체 회신용
   archive/<기준일자>/     처리 완료 파일

 ※ 환경변수(TLOG_HOME, TLOG_DB_URL, TLOG_DB_USER, TLOG_DB_PASSWORD 등)는
   서버 공통 프로파일에 설정되어 있음. 인프라팀 관리 사항이므로 임의 변경 금지.

3. 정기 실행 (crontab -l , tlogadm)
 # TLOG 배치 - 기준일자 = 전일
 30 6 * * * bash -lc '/app/tlog/batch/bin/run_job.sh LOAD  $(date -d "1 day ago" +\%Y\%m\%d)' >> /app/tlog/logs/load.log 2>&1
 40 6 * * * bash -lc '/app/tlog/batch/bin/run_job.sh JUDGE $(date -d "1 day ago" +\%Y\%m\%d)' >> /app/tlog/logs/judge.log 2>&1
 50 6 * * * bash -lc '/app/tlog/batch/bin/run_job.sh YIELD $(date -d "1 day ago" +\%Y\%m\%d)' >> /app/tlog/logs/yield.log 2>&1
 # 00 7 * * * ... ALERT  (2017.04 YIELD 잡에 통합되어 삭제)

 - MES 가 07:30 에 outbox 의 YIELD_<기준일자>.txt 를 수거한다.
   07:30 이전에 YIELD 가 끝나야 함. 늦어지면 MES팀에 통보.
 - 품질팀 메일 스크립트가 07:40 에 ALERT_<기준일자>.csv 가 있으면 경보 메일을 발송한다.
 - outbox 파일은 수작업으로 수정하지 말 것 (MES 수신 오류 발생 이력 있음).

4. 수동 실행 / 재실행
 - 실행 : /app/tlog/batch/bin/run_job.sh <LOAD|JUDGE|YIELD|PURGE> <yyyyMMdd>
 - 종료코드 0 = 정상, 그 외 = 실패 → 로그 확인
 - 동일 기준일자 재실행 가능 (스크립트가 run.ts 파라미터를 자동으로 붙임)
 - JUDGE, YIELD 는 재실행하면 해당 일자 결과를 다시 만든다 (중복 생성 없음)
 - LOAD 재처리
     처리된 파일은 archive/<기준일자>/ 로 이동되므로 재적재가 필요하면
     해당 파일을 inbox 로 복사한 뒤 LOAD 재실행.
     (이미 적재된 데이터와 키가 같은 레코드는 중복 건수로 집계됨)
 - 야간조(00~06시) 테스트는 근무일자가 전일이지만 익일자 파일로 들어오는 경우가 있음.
   현업에서 전일자 수율 재집계 요청이 오면 익일 LOAD 이후 아래 순서로 수행
     run_job.sh JUDGE <전일자>
     run_job.sh YIELD <전일자>

5. 장애 대응
 - LOAD 실패 : 로그에서 오류 파일 확인. 파일 깨짐/전송 중단이면 장비파트에 재전송 요청.
 - reject 발생 : 매일 오전 reject 디렉토리 확인 후 장비 업체 회신.
   파일 전체 오류(E09)는 해당 파일이 전혀 적재되지 않은 것이므로 즉시 재전송 요청.
 - YIELD 리포트 미생성 : MES팀에 지연 통보 후 재실행.
 - DB 접속 오류 : DBA 연락.
 - 기준정보(제품/BIN 룰) 오류 : 기준정보팀에 확인 요청. 배치에서 직접 수정 금지.

6. 연락처
 제조IT팀 TLOG 담당      이동욱 책임   (내선 4417)
 제조IT팀 (부담당)        최민호 선임   (내선 4421)
 DBA                    정보시스템팀 DB파트 (내선 2200)
 MES                    MES팀 한지민 (내선 3315)
 테스터 장비 (ADV/TER)    테스트기술팀 장비파트 (내선 5120)
 기준정보                기준정보팀 (기준정보 변경요청서)

7. 참고 문서
 - TLOG_배치_설계서_v1.3.docx
 - 인터페이스_정의서.xlsx
"""


def ops_txt():
    (OUT / "운영_가이드.txt").write_text(OPS, encoding="utf-8")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    design_doc()
    interface_xlsx()
    ops_txt()
    print("materials written to", OUT)
