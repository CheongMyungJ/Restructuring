"""Offline tests for rekit's deterministic parts (no AI calls). Run: python3 -m unittest discover -s tests"""
import json, os, re, shutil, subprocess, sys, tempfile, unittest

ENGINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ENGINE)
from rekit import inventory, materials  # noqa: E402
from rekit.questions import QuestionStore  # noqa: E402
from rekit.pipeline import Pipeline  # noqa: E402

RK = [sys.executable, os.path.join(ENGINE, 'rekit.py')]


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix='rekit_test_')
        self.code = os.path.join(self.tmp, 'code')
        os.makedirs(os.path.join(self.code, 'src'))
        open(os.path.join(self.code, 'src', 'ftl.c'), 'w').write(
            '#include "ftl.h"\nstatic int victim(int a) { return a; }\nint gc_run(void)\n{\n  return victim(1);\n}\n')
        open(os.path.join(self.code, 'src', 'ftl.h'), 'w').write('int gc_run(void);\n')
        open(os.path.join(self.code, 'src', 'Job.java'), 'w').write(
            'public class LoadJob {\n  public void execute(String d) {\n  }\n}\n')
        open(os.path.join(self.code, 'mapper.xml'), 'w').write('<mapper><select id="selectDevices">SELECT 1</select></mapper>')
        open(os.path.join(self.code, 'Makefile'), 'w').write('all:\n\tcc src/ftl.c\n')
        open(os.path.join(self.code, 'legacy.txt'), 'wb').write('한글 주석'.encode('cp949'))
        open(os.path.join(self.code, 'blob.bin'), 'wb').write(b'\x00\x01\x02')
        self.mat = os.path.join(self.tmp, 'docs')
        os.makedirs(self.mat)
        open(os.path.join(self.mat, 'note.md'), 'w').write('# 설계\n내용')

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_inventory(self):
        inv = inventory.scan([('code', self.code)])
        by = {f['path']: f for f in inv['files']}
        self.assertEqual(by['src/ftl.c']['lang'], 'c')
        self.assertIn('gc_run', by['src/ftl.c']['symbols'])
        self.assertIn('victim', by['src/ftl.c']['symbols'])
        self.assertIn('LoadJob', by['src/Job.java']['symbols'])
        self.assertIn('selectDevices', by['mapper.xml']['symbols'])
        self.assertEqual(by['blob.bin']['kind'], 'binary')
        self.assertEqual(by['legacy.txt']['encoding'], 'cp949')
        self.assertTrue(by['Makefile']['build_file'])

    def test_materials(self):
        idx = materials.convert_all([('materials', self.mat)], os.path.join(self.tmp, 'mt'))
        self.assertEqual(len(idx), 1)
        self.assertTrue(os.path.exists(os.path.join(self.tmp, 'mt', idx[0]['text'])))
        try:
            import docx  # noqa: F401
            import openpyxl
        except ImportError:
            return
        import docx as d
        doc = d.Document(); doc.add_heading('제목', 1); doc.add_paragraph('본문'); t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = 'A'; t.cell(1, 1).text = 'B'
        doc.save(os.path.join(self.mat, 'spec.docx'))
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Layout'; ws.append(['offset', 'len']); ws.append([0, 16])
        wb.save(os.path.join(self.mat, 'if.xlsx'))
        idx = materials.convert_all([('materials', self.mat)], os.path.join(self.tmp, 'mt2'))
        texts = {e['source']: open(os.path.join(self.tmp, 'mt2', e['text']), encoding='utf-8').read() for e in idx if e.get('text')}
        self.assertIn('# 제목', texts['materials:spec.docx'])
        self.assertIn('| A |', texts['materials:spec.docx'])
        self.assertIn('Sheet: Layout', texts['materials:if.xlsx'])
        self.assertIn('| 0 | 16 |', texts['materials:if.xlsx'])

    def test_question_roundtrip(self):
        ws = os.path.join(self.tmp, 'ws')
        os.makedirs(os.path.join(ws, '.rekit', 'questions'))
        qf = os.path.join(ws, '.rekit', 'questions', 'x.json')
        json.dump([{'id': 'Q-A-01', 'priority': 'blocking', 'question': '운영 시간대는?', 'options': ['KST', 'UTC']},
                   {'id': 'Q-A-02', 'priority': 'minor', 'question': '사용 안 하는 잡?'}], open(qf, 'w'), ensure_ascii=False)
        s = QuestionStore(ws)
        self.assertEqual(s.ingest_file(qf, 'A'), 2)
        self.assertEqual(s.ingest_file(qf, 'A'), 0)          # idempotent
        self.assertEqual([q['id'] for q in s.pending()], ['Q-A-01', 'Q-A-02'])
        out = os.path.join(self.tmp, 'q.md')
        self.assertEqual(s.export_for_humans(out), 2)
        txt = open(out, encoding='utf-8').read()
        txt = txt.replace('답변:\n\n\n## Q-A-02', '답변: 1\n\n\n## Q-A-02', 1).rstrip() + ' 보류\n'
        txt = txt.replace('답변: 보류', '답변: 보류')
        open(out, 'w', encoding='utf-8').write(txt)
        self.assertEqual(s.import_from_humans(out), 2)
        s.save()
        by = QuestionStore(ws).by_id()
        self.assertEqual(by['Q-A-01']['status'], 'answered')
        self.assertEqual(by['Q-A-01']['answer'], 'KST')        # option number resolved
        self.assertEqual(by['Q-A-02']['status'], 'deferred')
        self.assertTrue(os.path.exists(os.path.join(ws, 'kit', '09_open_questions.md')))

    def test_cli_init_inventory_status(self):
        ws = os.path.join(self.tmp, 'work')
        r = subprocess.run(RK + ['init', ws, '--code', self.code, '--materials', self.mat, '--no-intake',
                                 '--intake-json', '{"goal": "플랫폼 교체"}'], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn('플랫폼 교체', open(os.path.join(ws, '.rekit', 'intake.md'), encoding='utf-8').read())
        r = subprocess.run(RK + ['run', ws, '--until', 'inventory'], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertTrue(os.path.exists(os.path.join(ws, '.rekit', 'inventory.json')))
        self.assertTrue(os.path.exists(os.path.join(ws, '.rekit', 'prompts', 'kitspec.md')))
        self.assertNotIn('{{LANG}}', open(os.path.join(ws, '.rekit', 'prompts', 'kitspec.md'), encoding='utf-8').read())
        r = subprocess.run(RK + ['run', ws, '--until', 'inventory'], capture_output=True, text=True, timeout=60)
        self.assertEqual(r.returncode, 0, r.stderr)          # already done: must not proceed to survey
        self.assertIn('일시정지', r.stdout)
        self.assertFalse(os.path.exists(os.path.join(ws, '.rekit', 'logs')))
        r = subprocess.run(RK + ['status', ws, '--json'], capture_output=True, text=True)
        st = json.loads(r.stdout.strip().splitlines()[-1])
        self.assertEqual(st['phases']['inventory'], 'done')
        self.assertEqual(st['run_state'], 'paused')           # only inventory ran
        self.assertFalse(os.path.exists(os.path.join(ws, '.rekit', 'run.pid')))

    def test_coverage(self):
        ws = os.path.join(self.tmp, 'work2')
        subprocess.run(RK + ['init', ws, '--code', self.code, '--no-intake'], capture_output=True, text=True, check=True)
        subprocess.run(RK + ['run', ws, '--until', 'inventory'], capture_output=True, text=True, check=True)
        os.makedirs(os.path.join(ws, 'kit', '05_behavior'), exist_ok=True)
        open(os.path.join(ws, 'kit', '05_behavior', 'GC.md'), 'w').write('gc_run (src: src/ftl.c:3-6)')
        cov = Pipeline(ws).coverage()
        self.assertIn('code:src/ftl.c#victim', cov['unreferenced_symbols'])
        self.assertNotIn('code:src/ftl.c#gc_run', cov['unreferenced_symbols'])
        self.assertNotIn('code:src/ftl.c', cov['unreferenced_files'])

    def test_state_concurrency(self):
        import threading
        ws = os.path.join(self.tmp, 'work4')
        subprocess.run(RK + ['init', ws, '--code', self.code, '--no-intake'], capture_output=True, text=True, check=True)
        p = Pipeline(ws)

        def worker(i):
            for j in range(50):
                p.update_state(lambda s: s['sessions'].append({'step': f'w{i}', 'cost_usd': 0.01}))
                p.update_state(lambda s: s['units'].__setitem__(f'U{i}_{j}', 'done'))

        ths = [threading.Thread(target=worker, args=(i,)) for i in range(8)]
        [t.start() for t in ths]; [t.join() for t in ths]
        st = json.load(open(os.path.join(ws, '.rekit', 'state.json')))
        self.assertEqual(len(st['sessions']), 400)
        self.assertEqual(len(st['units']), 400)
        self.assertAlmostEqual(Pipeline(ws).total_cost(), 4.0, places=2)

    def test_failed_session_not_marked_done(self):
        from rekit import pipeline as pl
        ws = os.path.join(self.tmp, 'work5')
        subprocess.run(RK + ['init', ws, '--code', self.code, '--no-intake'], capture_output=True, text=True, check=True)
        fake = {'result': "You've hit your session limit", 'structured': None, 'cost_usd': 0.0, 'wall_s': 0.1,
                'num_turns': 1, 'session_id': 'x', 'is_error': True, 'timed_out': False, 'usage_limit': True}
        orig = pl.run_claude
        pl.run_claude = lambda *a, **k: dict(fake)
        try:
            p = Pipeline(ws, log=lambda *a: None)
            with self.assertRaises(pl.SessionFailed):
                p.run()
            st = json.load(open(os.path.join(ws, '.rekit', 'state.json')))
            self.assertEqual(st['phases']['inventory'], 'done')
            self.assertNotEqual(st['phases'].get('survey'), 'done')
        finally:
            pl.run_claude = orig

    def test_prompts_render(self):
        ws = os.path.join(self.tmp, 'work3')
        subprocess.run(RK + ['init', ws, '--code', self.code, '--no-intake'], capture_output=True, text=True, check=True)
        p = Pipeline(ws)
        for name, kw in (('survey.md', {'UNIT_TOKENS': 1}), ('unit_v2.md', {'UNIT_ID': 'GC', 'UNIT_TITLE': 't', 'FILES': '-', 'FOCUS': '', 'CHECKLIST': 'x', 'CAPS': 'CAP-SPACE'}),
                         ('specify.md', {'CAP_ID': 'CAP-SPACE', 'CAP_TITLE': 't', 'AREA': 'SPACE', 'CAP_DESC': 'd', 'NOTES': 'n', 'ALIASES': ''}),
                         ('neutralize.md', {}), ('design.md', {}), ('contracts.md', {}),
                         ('synthesize_v2.md', {'SOURCES': 's'}), ('synthesize_group.md', {'GROUP_ID': 'G', 'GROUP_TITLE': 't', 'UNITS': 'GC'}),
                         ('characterize.md', {}), ('integrate.md', {}), ('review_v2.md', {'AREA': 'GC', 'AREA_ID': 'GC', 'FILES': 'f'}),
                         ('fix.md', {'BATCH': 'GC'}), ('finalize.md', {})):
            out = p.render(name, **kw)
            self.assertNotIn('{{', out, name)


    def _ws_after_inventory(self, name):
        ws = os.path.join(self.tmp, name)
        subprocess.run(RK + ['init', ws, '--code', self.code, '--no-intake'], capture_output=True, text=True, check=True)
        subprocess.run(RK + ['run', ws, '--until', 'inventory'], capture_output=True, text=True, check=True)
        return ws

    def test_neutrality_scan(self):
        from rekit import neutrality
        ws = self._ws_after_inventory('work6')
        kit = os.path.join(ws, 'kit')
        for d in ('05_behavior', '03_interfaces', 'legacy_ref/modules'):
            os.makedirs(os.path.join(kit, d), exist_ok=True)
        open(os.path.join(kit, '05_behavior', 'CAP-SPACE.md'), 'w').write(
            '### BR-SPACE-001 x\n`gc_run` picks the victim (src: src/ftl.c:3-6), see ftl.c and LoadJob.\n')
        open(os.path.join(kit, 'legacy_ref', 'modules', 'GC.md'), 'w').write('gc_run (src: src/ftl.c:3-6)')
        r = neutrality.scan(ws)
        f = r['files']['05_behavior/CAP-SPACE.md']
        self.assertEqual(f['symbols'], {'gc_run': 1, 'LoadJob': 1})   # 'victim' is a plain word: not code-like
        self.assertEqual(f['file_names'], {'ftl.c': 1})
        self.assertEqual(f['src_refs'], 1)
        self.assertNotIn('legacy_ref/modules/GC.md', r['files'])        # legacy_ref may name everything
        open(os.path.join(kit, '03_interfaces', 'IF-API-001.md'), 'w').write('int gc_run(void);')
        r = neutrality.scan(ws)                                          # contract identifiers are allowed
        self.assertNotIn('gc_run', r['files']['05_behavior/CAP-SPACE.md']['symbols'])

    def test_specify_and_clean_room_design(self):
        from rekit import pipeline as pl
        ws = self._ws_after_inventory('work7')
        kit, rk = os.path.join(ws, 'kit'), os.path.join(ws, '.rekit')
        json.dump({'units': [{'id': 'GC', 'files': ['code:src/ftl.c'], 'capabilities': ['CAP-SPACE']}],
                   'capabilities': [{'id': 'CAP-SPACE', 'area': 'SPACE', 'title': 'space', 'units': ['GC']},
                                    {'id': 'CAP-HEALTH', 'area': 'HEALTH', 'title': 'health', 'units': []}]},
                  open(os.path.join(rk, 'units.json'), 'w'))
        os.makedirs(os.path.join(rk, 'notes', 'GC'))
        open(os.path.join(rk, 'notes', 'GC', 'behavior.md'), 'w').write('## CAP-SPACE\n### BR-GC-001\n## CAP-OLDH\n')
        json.dump({'LR-GC-001': {'src': ['code:src/ftl.c:2'], 'confidence': 'x'}, 'BR-GC-001': {'src': []}},
                  open(os.path.join(rk, 'notes', 'GC', 'trace.json'), 'w'))
        json.dump({'CAP-OLDH': 'CAP-HEALTH'}, open(os.path.join(rk, 'capability_map.json'), 'w'))
        os.makedirs(os.path.join(kit, 'legacy_ref'), exist_ok=True)
        open(os.path.join(kit, 'legacy_ref', 'structure_issues.md'), 'w').write('### LP-001 global state')
        open(os.path.join(kit, '00_overview.md'), 'w').write('overview')
        calls = []

        def fake(prompt, cwd, log_path, **k):
            calls.append({'prompt': prompt, 'cwd': cwd, **k})
            if 'Phase: SPECIFY' in prompt:
                cap = re.search(r'capability (CAP-\w+)', prompt).group(1)
                os.makedirs(os.path.join(cwd, 'kit', '05_behavior'), exist_ok=True)
                open(os.path.join(cwd, 'kit', '05_behavior', cap + '.md'), 'w').write('spec')
                os.makedirs(os.path.join(cwd, '.rekit', 'notes', 'spec'), exist_ok=True)
                json.dump({f'BR-{cap[4:]}-001': {'file': 'kit/05_behavior/' + cap + '.md', 'src': ['code:src/ftl.c:3']}},
                          open(os.path.join(cwd, '.rekit', 'notes', 'spec', cap + '.trace.json'), 'w'))
            if 'TARGET DESIGN' in prompt:
                os.makedirs(os.path.join(cwd, 'design'))
                open(os.path.join(cwd, 'design', 'architecture.md'), 'w').write(
                    'saw legacy_ref: %s, issues: %s' % (os.path.exists(os.path.join(cwd, 'legacy_ref')),
                                                        open(os.path.join(cwd, 'inputs', 'structure_issues.md')).read()))
            return {'result': 'ok', 'structured': None, 'cost_usd': 0.01, 'wall_s': 0.1, 'num_turns': 1,
                    'session_id': 'x', 'is_error': False, 'timed_out': False}
        orig = pl.run_claude
        pl.run_claude = fake
        try:
            p = Pipeline(ws, log=lambda *a: None)
            p.p_specify()
            p.p_design()
        finally:
            pl.run_claude = orig
        health = next(c for c in calls if 'capability CAP-HEALTH' in c['prompt'])
        self.assertIn('.rekit/notes/GC/behavior.md', health['prompt'])      # found through the alias heading
        self.assertIn('`## CAP-OLDH`', health['prompt'])
        trace = json.load(open(os.path.join(kit, 'trace.json')))
        self.assertIn('LR-GC-001', trace)
        self.assertNotIn('BR-GC-001', trace)                                  # unit note ids are not kit ids
        self.assertEqual(trace['BR-SPACE-001']['src'], ['code:src/ftl.c:3'])
        d = calls[-1]
        self.assertEqual(d['add_dirs'], [])                                   # clean room: no legacy code access
        self.assertNotIn('Bash', d['allowed_tools'])
        self.assertEqual(open(os.path.join(kit, 'design', 'architecture.md')).read(),
                         'saw legacy_ref: False, issues: ### LP-001 global state')
        self.assertFalse(os.path.exists(os.path.join(rk, 'design_box')))


if __name__ == '__main__':
    unittest.main()
