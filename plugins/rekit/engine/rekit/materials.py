"""Convert handed-over materials (PDF, Office, HWP(X), HTML, text) into text files the AI can read efficiently."""
import json, os, re, shutil, subprocess, zipfile

TEXT_EXT = {'.txt', '.md', '.csv', '.tsv', '.sql', '.json', '.xml', '.yaml', '.yml', '.properties', '.ini', '.log', '.rst'}
IMAGE_EXT = {'.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp'}


def _run(cmd, timeout=300):
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        return p.returncode, p.stdout
    except Exception as e:  # missing tool / timeout
        return 1, str(e)


def pdf_to_text(path):
    if shutil.which('pdftotext'):
        rc, out = _run(['pdftotext', '-layout', path, '-'])
        if rc == 0 and out.strip():
            pages = out.split('\f')
            return '\n'.join(f'\n--- page {i + 1} ---\n{p}' for i, p in enumerate(pages) if p.strip())
    try:
        import pypdf
        r = pypdf.PdfReader(path)
        return '\n'.join(f'\n--- page {i + 1} ---\n{(pg.extract_text() or "")}' for i, pg in enumerate(r.pages))
    except Exception as e:
        return f'[PDF text extraction failed: {e}]'


def docx_to_text(path):
    try:
        import docx
        d = docx.Document(path)
        out = []
        for block in d.element.body.iterchildren():
            tag = block.tag.split('}')[-1]
            if tag == 'p':
                p = docx.text.paragraph.Paragraph(block, d)
                style = (p.style.name or '').lower() if p.style is not None else ''
                txt = p.text
                if style.startswith('heading'):
                    lvl = re.findall(r'\d+', style)
                    txt = '#' * (int(lvl[0]) if lvl else 1) + ' ' + txt
                out.append(txt)
            elif tag == 'tbl':
                t = docx.table.Table(block, d)
                for i, row in enumerate(t.rows):
                    cells = [c.text.replace('\n', ' ').strip() for c in row.cells]
                    out.append('| ' + ' | '.join(cells) + ' |')
                    if i == 0:
                        out.append('|' + '---|' * len(cells))
                out.append('')
        return '\n'.join(out)
    except Exception:
        if shutil.which('pandoc'):
            rc, out = _run(['pandoc', path, '-t', 'gfm'])
            if rc == 0:
                return out
        return '[DOCX conversion failed]'


def xlsx_to_text(path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f'\n## Sheet: {ws.title}\n')
            rows = [[('' if v is None else str(v)).replace('\n', ' ').replace('|', '\\|') for v in r]
                    for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(c.strip() for c in r)]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            while width > 1 and all((len(r) < width or not r[width - 1].strip()) for r in rows):
                width -= 1
            for i, r in enumerate(rows):
                r = (r + [''] * width)[:width]
                out.append('| ' + ' | '.join(r) + ' |')
                if i == 0:
                    out.append('|' + '---|' * width)
        return '\n'.join(out)
    except Exception as e:
        return f'[XLSX conversion failed: {e}]'


def pptx_to_text(path):
    try:
        import pptx
        prs = pptx.Presentation(path)
        out = []
        for i, slide in enumerate(prs.slides, 1):
            out.append(f'\n--- slide {i} ---')
            for shp in slide.shapes:
                if shp.has_text_frame:
                    out.append(shp.text_frame.text)
                if getattr(shp, 'has_table', False) and shp.has_table:
                    for row in shp.table.rows:
                        out.append('| ' + ' | '.join(c.text for c in row.cells) + ' |')
        return '\n'.join(out)
    except Exception as e:
        return f'[PPTX conversion failed: {e}]'


def hwpx_to_text(path):
    try:
        z = zipfile.ZipFile(path)
        parts = sorted(n for n in z.namelist() if re.match(r'Contents/section\d+\.xml', n))
        out = []
        for n in parts:
            xml = z.read(n).decode('utf-8', 'replace')
            xml = re.sub(r'</hp:p>', '\n', xml)
            out.append(re.sub(r'<[^>]+>', '', xml))
        return '\n'.join(out)
    except Exception as e:
        return f'[HWPX conversion failed: {e}]'


def via_libreoffice(path, fmt):
    lo = shutil.which('soffice') or shutil.which('libreoffice')
    if not lo:
        return None
    outdir = os.path.join(os.path.dirname(path), '.lo_tmp')
    os.makedirs(outdir, exist_ok=True)
    rc, _ = _run([lo, '--headless', '--convert-to', fmt, '--outdir', outdir, path], timeout=600)
    cands = [os.path.join(outdir, f) for f in os.listdir(outdir)]
    return cands[0] if cands else None


def convert_one(path):
    """Return (text or None, note). None text means: keep original (e.g. images) for direct viewing."""
    ext = os.path.splitext(path)[1].lower()
    if ext in TEXT_EXT:
        with open(path, 'rb') as f:
            raw = f.read()
        for enc in ('utf-8-sig', 'cp949', 'latin-1'):
            try:
                return raw.decode(enc), f'text ({enc})'
            except UnicodeDecodeError:
                pass
    if ext == '.pdf':
        return pdf_to_text(path), 'pdf'
    if ext == '.docx':
        return docx_to_text(path), 'docx'
    if ext in ('.xlsx', '.xlsm'):
        return xlsx_to_text(path), 'xlsx'
    if ext == '.pptx':
        return pptx_to_text(path), 'pptx'
    if ext == '.hwpx':
        return hwpx_to_text(path), 'hwpx'
    if ext == '.hwp':
        if shutil.which('hwp5txt'):
            rc, out = _run(['hwp5txt', path])
            if rc == 0:
                return out, 'hwp (hwp5txt)'
        conv = via_libreoffice(path, 'docx')
        if conv:
            return docx_to_text(conv), 'hwp (libreoffice)'
        return None, 'hwp: no converter available (install pyhwp: pip install pyhwp) — ask user for PDF export'
    if ext in ('.doc', '.rtf', '.odt'):
        conv = via_libreoffice(path, 'docx')
        return (docx_to_text(conv), ext[1:] + ' (libreoffice)') if conv else (None, 'no converter')
    if ext in ('.xls', '.ods'):
        conv = via_libreoffice(path, 'xlsx')
        return (xlsx_to_text(conv), ext[1:] + ' (libreoffice)') if conv else (None, 'no converter')
    if ext in ('.ppt', '.odp'):
        conv = via_libreoffice(path, 'pptx')
        return (pptx_to_text(conv), ext[1:] + ' (libreoffice)') if conv else (None, 'no converter')
    if ext in ('.html', '.htm'):
        with open(path, 'rb') as f:
            txt = f.read().decode('utf-8', 'replace')
        return re.sub(r'\s+\n', '\n', re.sub(r'<[^>]+>', ' ', txt)), 'html'
    if ext in IMAGE_EXT:
        return None, 'image (view directly)'
    return None, 'unsupported (view directly if possible)'


def convert_all(material_roots, out_dir):
    """material_roots: list of (label, abs_path). Writes <out_dir>/<label>/<relpath>.md; returns index list."""
    os.makedirs(out_dir, exist_ok=True)
    index = []
    for label, root in material_roots:
        paths = [root] if os.path.isfile(root) else [
            os.path.join(dp, f) for dp, dn, fs in os.walk(root) for f in sorted(fs) if not dp.endswith('.lo_tmp')]
        for p in sorted(paths):
            rel = os.path.basename(p) if os.path.isfile(root) else os.path.relpath(p, root)
            text, note = convert_one(p)
            entry = {'source': f'{label}:{rel}', 'abs': p, 'note': note, 'bytes': os.path.getsize(p)}
            if text is not None:
                dst = os.path.join(out_dir, label, rel + '.md')
                os.makedirs(os.path.dirname(dst), exist_ok=True)
                with open(dst, 'w', encoding='utf-8') as f:
                    f.write(f'<!-- converted from {label}:{rel} ({note}) -->\n' + text)
                entry.update(text=os.path.relpath(dst, out_dir), words=len(text.split()))
            index.append(entry)
    with open(os.path.join(out_dir, 'index.json'), 'w', encoding='utf-8') as f:
        json.dump(index, f, ensure_ascii=False, indent=1)
    return index
